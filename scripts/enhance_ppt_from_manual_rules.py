from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


SRC = Path("docs/yinren_kg_agent_source.pptx")
OUT = Path("docs/yinren_kg_agent_presentation_enhanced.pptx")

BLUE = RGBColor(47, 101, 166)
DARK = RGBColor(33, 37, 41)
MUTED = RGBColor(88, 96, 105)
LIGHT_BG = RGBColor(247, 249, 252)
ACCENT = RGBColor(200, 80, 60)


def clear_slide(slide):
    sp_tree = slide.shapes._spTree
    for shape in list(slide.shapes):
        sp_tree.remove(shape._element)


def set_run(run, size=18, bold=False, color=DARK):
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_box(slide, left, top, width, height, text="", fill=None, line=None):
    shape = slide.shapes.add_shape(1, left, top, width, height)
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.color.rgb = RGBColor(230, 234, 240)
    else:
        shape.line.color.rgb = line
    if text:
        shape.text = text
    return shape


def add_text(slide, left, top, width, height, text, size=18, bold=False, color=DARK, align=None):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    if align is not None:
        p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, color=color)
    return box


def add_bullets(slide, left, top, width, height, items, size=16, color=DARK, gap=0):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.clear()
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.space_after = Pt(gap)
        for run in p.runs:
            set_run(run, size=size, color=color)
    return box


def make_title_slide(slide, title, subtitle):
    clear_slide(slide)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(250, 251, 253)
    add_text(slide, Inches(0.75), Inches(1.7), Inches(11.8), Inches(0.7), title, size=34, bold=True, color=BLUE, align=PP_ALIGN.CENTER)
    add_text(slide, Inches(1.35), Inches(2.55), Inches(10.6), Inches(0.5), subtitle, size=18, color=MUTED, align=PP_ALIGN.CENTER)
    add_box(slide, Inches(1.1), Inches(3.45), Inches(10.95), Inches(1.55), fill=LIGHT_BG)
    add_bullets(
        slide,
        Inches(1.45),
        Inches(3.65),
        Inches(10.2),
        Inches(1.15),
        [
            "课程技能点：知识抽取、RDF/Turtle、本体、实体链接、LangGraph、工具调用、few-shot SPARQL、网页前端",
            "展示重点：功能展示 + 实现解释 + 人工设计规则与 AI 辅助边界",
        ],
        size=16,
    )


def make_two_col(slide, title, left_title, left_items, right_title, right_items, footer=None):
    clear_slide(slide)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
    add_text(slide, Inches(0.55), Inches(0.35), Inches(12.2), Inches(0.45), title, size=24, bold=True, color=BLUE)
    add_box(slide, Inches(0.65), Inches(1.1), Inches(5.95), Inches(5.2), fill=LIGHT_BG)
    add_box(slide, Inches(6.85), Inches(1.1), Inches(5.95), Inches(5.2), fill=LIGHT_BG)
    add_text(slide, Inches(0.95), Inches(1.35), Inches(5.35), Inches(0.35), left_title, size=17, bold=True, color=BLUE)
    add_bullets(slide, Inches(0.95), Inches(1.85), Inches(5.35), Inches(3.95), left_items, size=14)
    add_text(slide, Inches(7.15), Inches(1.35), Inches(5.35), Inches(0.35), right_title, size=17, bold=True, color=BLUE)
    add_bullets(slide, Inches(7.15), Inches(1.85), Inches(5.35), Inches(3.95), right_items, size=14)
    if footer:
        add_text(slide, Inches(0.75), Inches(6.55), Inches(12), Inches(0.35), footer, size=12, color=MUTED)


def make_flow(slide, title, steps, note):
    clear_slide(slide)
    add_text(slide, Inches(0.55), Inches(0.35), Inches(12), Inches(0.45), title, size=24, bold=True, color=BLUE)
    x = Inches(0.75)
    y = Inches(1.55)
    w = Inches(2.0)
    h = Inches(0.8)
    for i, step in enumerate(steps):
        add_box(slide, x + Inches(2.1 * i), y, w, h, fill=LIGHT_BG, line=BLUE)
        add_text(slide, x + Inches(2.1 * i) + Inches(0.1), y + Inches(0.15), w - Inches(0.2), h - Inches(0.2), step, size=13, bold=True, color=DARK, align=PP_ALIGN.CENTER)
        if i < len(steps) - 1:
            add_text(slide, x + Inches(2.1 * i) + Inches(1.95), y + Inches(0.22), Inches(0.3), Inches(0.3), "->", size=18, bold=True, color=BLUE)
    add_box(slide, Inches(0.9), Inches(3.05), Inches(11.6), Inches(2.45), fill=LIGHT_BG)
    add_bullets(slide, Inches(1.2), Inches(3.35), Inches(11), Inches(1.85), note, size=15)


def delete_slide(prs, slide):
    slide_id = slide.slide_id
    slide_id_list = prs.slides._sldIdLst
    for sld_id in list(slide_id_list):
        if int(sld_id.get("id")) == slide_id:
            r_id = sld_id.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
            prs.part.drop_rel(r_id)
            slide_id_list.remove(sld_id)
            break


def add_manual_ai_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    make_two_col(
        slide,
        "人工设计与 AI 辅助边界",
        "人工设计 / 人工指导完成",
        [
            "抽取 Schema：实体类型、关系类型、输出结构",
            "数据清洗规则：去重、停用词、长度限制、中文字符过滤",
            "本体设计：9 类、18 个对象属性、7 个数据属性、domain/range",
            "实体链接规则：字号、籍贯、朝代、生卒年、多候选优先级",
            "LangGraph 路由：工具优先、SPARQL 生成、fallback",
            "few-shot SPARQL：典型问句与可执行模板",
        ],
        "AI 辅助完成",
        [
            "根据人物传记文本抽取实体、属性和关系",
            "复杂问题下辅助生成 SPARQL 查询",
            "根据 SPARQL / 工具结果组织自然语言回答",
            "辅助代码实现、页面优化和文档整理",
            "注意：LLM 打分消歧当前作为后续补充，不说成已完成",
        ],
        footer="核心答辩口径：不是单纯调用大模型，而是用人工规则约束 AI，使其服务于本地 RDF 知识图谱查询。",
    )


def add_fewshot_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    make_two_col(
        slide,
        "few-shot SPARQL：让模型生成可执行查询",
        "人工准备的典型问题",
        [
            "文彭的字是什么？",
            "文彭的号是什么？",
            "文彭的字和号是什么？",
            "文彭的生卒年是什么？",
            "文徵明与文彭是什么关系？",
            "谁开创了吴门印派？",
        ],
        "约束模型生成的关键点",
        [
            "统一 yrz 命名空间和 RDF/Turtle 本体",
            "兼容 core.ttl 的关系实例模型：Relation + relationType",
            "兼容 aligned.ttl 中外部补充的字面量年份",
            "复杂查询先生成 SPARQL，再执行并基于结果回答",
            "文件依据：src/fewshot_sparql.py、data/examples/fewshot_sparql.md",
        ],
    )


def add_demo_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    make_flow(
        slide,
        "建议演示路线",
        ["智能问答", "高级 SPARQL", "关系网络图", "图结构分析", "本体解释"],
        [
            "智能问答：输入“文彭的字和号是什么？”展示工具调用 / SPARQL / 回答。",
            "高级 SPARQL：查询 Person 或师承关系，证明 RDF/Turtle 图谱可被直接检索。",
            "关系网络图：展示人物节点与师承、亲属、交游、流派等边。",
            "图结构分析：展示中心性、社区、路径或流派分析，说明分析基于人物关系网络。",
            "本体解释：点击边展示实例关系、本体属性、domain、range 和中文说明。",
        ],
    )


def main():
    prs = Presentation(SRC)

    make_title_slide(
        prs.slides[0],
        "基于《印人传》的知识图谱智能体",
        "功能展示 + 课程技能点实现解释 + 人工规则与 AI 辅助边界",
    )

    make_two_col(
        prs.slides[1],
        "目录 CONTENTS：按课程采分点组织",
        "基础构建",
        [
            "01 系统总览：从文本到知识图谱智能体",
            "02 OCR / 文本整理与数据清洗",
            "03 知识抽取：人工 Schema + AI 辅助抽取",
            "04 关系构建：规则抽取 + 人物关系网络",
            "05 RDF/Turtle 与本体设计",
        ],
        "应用与展示",
        [
            "06 实体链接：ctext / cbdb 对齐与规则消歧",
            "07 LangGraph 问答：工具 + SPARQL + fallback",
            "08 few-shot SPARQL：典型问句模板",
            "09 前端展示：问答、SPARQL、关系图、图分析、本体解释",
            "10 人工设计与 AI 辅助边界说明",
        ],
    )

    make_flow(
        prs.slides[2],
        "01 项目总架构：从古籍文本到可问答知识图谱",
        ["文本整理", "知识抽取", "RDF/Turtle", "实体链接", "LangGraph", "前端展示"],
        [
            "数据源：以《印人传》整理后人物记录为主，外部补充使用 ctext / cbdb。",
            "知识表示：schema.ttl 定义本体，core.ttl 存储核心事实，aligned.ttl 存储外部对齐结果。",
            "应用层：LangGraph 工作流优先调用固定工具，复杂问题转 few-shot SPARQL，失败时 fallback。",
            "前端：提供智能问答、高级 SPARQL、关系网络图、图结构分析和本体解释展示。",
        ],
    )

    make_two_col(
        prs.slides[3],
        "02 OCR / 文本整理与数据清洗",
        "当前可证明的数据结果",
        [
            "原始 PDF/OCR 后文本已整理为人物记录 JSON",
            "最终人物记录：1673 条",
            "保留 source_text / evidence，便于追溯原文",
            "仓库中保留清洗后的 JSON 数据；OCR 脚本当前未找到",
        ],
        "人工清洗规则",
        [
            "去重：seen_entities 避免重复实体",
            "停用词过滤：排除普通词误识别为实体",
            "长度限制：过滤异常人名、地名、字号",
            "中文字符过滤：减少 OCR 噪声",
            "原文窗口：保留抽取证据",
        ],
        footer="文件依据：data/source/person_records_v5.json、scripts/extraction/rule_based_extract.py",
    )

    make_two_col(
        prs.slides[5],
        "03 知识抽取：人工 Schema + AI 辅助抽取",
        "人工定义的实体类型",
        [
            "人物、地名、时间、字号",
            "书体印风、流派、印章、作品",
            "输出字段包含 name、type/category、source_text/evidence、confidence",
        ],
        "人工定义的关系类型",
        [
            "父子、师承、交游、流派归属、开创",
            "字号对应、籍贯、活动于、任职于",
            "生于、卒于、擅长、创作、著有",
            "AI 根据 Schema 抽取，规则脚本用于补充和兜底",
        ],
        footer="文件依据：data/intermediate/extraction_schema.json、src/information_extraction.py、scripts/extraction/rule_based_extract.py",
    )

    make_two_col(
        prs.slides[6],
        "03 知识抽取：关键数据与可追溯证据",
        "数据规模",
        [
            "人物记录：1673 条",
            "实体数据来自 data/intermediate/entities.json",
            "关系数据来自 data/intermediate/relations.json",
            "每条抽取结果尽量保留 source_text / evidence",
            "用于后续 Turtle 图谱构建和 SPARQL 查询",
        ],
        "证明方式",
        [
            "展示 extraction_schema.json：说明实体/关系类型是人工定义",
            "展示 rule_based_extract.py：说明清洗与规则抽取不是黑盒",
            "展示 entities.json / relations.json 片段：说明结果可追溯",
            "强调：AI 辅助抽取，人工规则约束输出边界",
        ],
        footer="谨慎口径：不要说仓库中保留了完整 OCR 脚本；当前可证明的是清洗后人物记录 JSON 和抽取结果。",
    )

    make_two_col(
        prs.slides[7],
        "04 关系构建：从文本规则到人物关系网络",
        "人工关系抽取规则",
        [
            "“字XX / 号XX / 别号XX” -> 字号关系",
            "“XX人 / 居XX / 家于XX” -> 籍贯或活动地",
            "“师XX / 从XX学” -> 师承关系",
            "“与XX游 / 同里XX” -> 交游关系",
            "“XX之子 / XX子 / XX孙” -> 亲属关系",
        ],
        "实际关系数据",
        [
            "关系实例：12498 条",
            "实际使用关系类型：15 种",
            "高频关系：字号、籍贯、技艺、著作",
            "网络关系：师承、交游、父子、流派",
            "低频高价值：foundsSchool 开创流派 11 条",
        ],
        footer="说明：原 PPT 中“19 类关系”偏宣传口径，实际 core.ttl 中使用 15 种 relationType。",
    )

    make_two_col(
        prs.slides[8],
        "04 关系构建：关键数据与示例",
        "实际关系类型分布",
        [
            "hasCourtesyName：2472",
            "birthPlace：2175",
            "hasArtName：1517",
            "skilledIn：1462",
            "authored：1166",
            "hasFriend：941",
            "hasTeacher：742",
            "fatherOf：425",
        ],
        "典型三元组讲法",
        [
            "文彭 -> hasCourtesyName -> 寿承",
            "文彭 -> hasArtName -> 三桥",
            "文彭 -> fatherOf / family relation -> 文徵明相关家族线索",
            "文彭 -> foundsSchool -> 吴门",
            "关系不是孤立展示，而是进入关系网络图和图结构分析",
        ],
        footer="答辩口径：关系类型数量不是越多越好，关键是能覆盖人物属性、师承、亲属、交游、流派等课程要求。",
    )

    make_two_col(
        prs.slides[9],
        "05 本体与 Turtle 存储：标准化知识组织",
        "人工设计的本体规模",
        [
            "核心类：9 个",
            "对象属性：18 个",
            "数据属性：7 个",
            "Person 实例：5344",
            "Relation 实例：12498",
        ],
        "核心类与属性示例",
        [
            "Person -> hasTeacher -> Person",
            "Person -> fatherOf -> Person",
            "Person -> hasFriend -> Person",
            "Person -> belongsToSchool -> School",
            "Person -> hasCourtesyName -> CourtesyName",
            "Person -> birthPlace -> Place",
            "Person -> skilledIn -> CalligraphyStyle",
        ],
        footer="文件依据：data/kg/schema.ttl、data/kg/core.ttl、data/kg/aligned.ttl；查询层通过 RDFLib 加载 Turtle。",
    )

    make_two_col(
        prs.slides[10],
        "05 本体解释如何在前端展示",
        "本体层",
        [
            "Class：Person、School、Place、TimePeriod、Work 等",
            "Property：hasTeacher、fatherOf、hasFriend、belongsToSchool 等",
            "domain/range：约束关系两端的实体类型",
            "schema.ttl 负责定义本体，core.ttl 负责实例事实",
        ],
        "实例层可视化",
        [
            "人物关系图展示的是实例节点与实例关系",
            "点击边时解释其本体属性",
            "示例：文彭 --师承--> 某人物，对应 yrz:hasTeacher",
            "展示 domain=Person，range=Person，说明该关系含义",
            "这比单纯文字介绍本体更适合答辩演示",
        ],
        footer="这页用于回答“本体构建有没有展示出来”：通过关系网络图点击解释本体层。",
    )

    make_two_col(
        prs.slides[11],
        "06 实体链接与知识补充：规则消歧优先",
        "当前对齐结果",
        [
            "核心人物优先对齐，不说全量完成",
            "唯一已对齐人物：20 个",
            "aligned 资源：104 个",
            "owl:sameAs：165 条",
            "CBDB ID 资源：104 个，CText ID 资源：61 个",
        ],
        "人工消歧规则",
        [
            "字号完全相同 + 籍贯相同：高置信",
            "字号相同 + 籍贯同区域：中置信",
            "仅字号相同：低置信，需人工确认",
            "朝代 / 生卒年不匹配：排除候选",
            "多候选：优先籍贯完全匹配，再看同省同府",
            "LLM 打分消歧：后续补充，不说已完成",
        ],
        footer="文件依据：data/kg/alignment_rules.md、data/kg/aligned.ttl",
    )

    make_two_col(
        prs.slides[12],
        "06 实体链接：文彭案例与谨慎口径",
        "文彭对齐示例",
        [
            "本地图谱人物：文彭",
            "CBDB ID：34677",
            "CText ID：326035",
            "补充信息：bornIn 1498，外部 sameAs 链接",
            "aligned.ttl 中保存 owl:sameAs、cbdbId、ctextId",
        ],
        "答辩说明",
        [
            "不是全量 5344 人都已确认对齐",
            "当前采用核心人物优先策略",
            "大量人物缺少字号、籍贯、生卒年等消歧线索",
            "规则无法唯一判断时，后续加入 LLM 打分消歧",
            "不要使用未核实的外部数据库 ID",
        ],
        footer="这页已替换旧 PPT 中不准确的未核实外部 ID 口径。",
    )

    make_two_col(
        prs.slides[13],
        "07 问答系统：LangGraph 三层链路",
        "工作流路径",
        [
            "用户问题进入 LangGraph 工作流",
            "LLM 判断是否调用固定工具",
            "ToolNode 执行工具并返回 ToolMessage",
            "工具有结果：LLM 根据结果组织回答",
            "工具无结果：转 few-shot / LLM 生成 SPARQL",
            "SPARQL 失败：进入 fallback 谨慎回答",
        ],
        "人工路由规则",
        [
            "字、号、字号 -> get_courtesy / get_art 工具",
            "生卒年 -> get_birth_death",
            "师承 -> get_teacher_relations",
            "亲属 -> get_family_relations",
            "交游 -> get_social_relations",
            "两人关系 -> get_pair_relations",
        ],
        footer="文件依据：src/tool_calling_workflow.py、src/workflow.py、src/langchain_tools.py、src/tools.py",
    )

    make_two_col(
        prs.slides[14],
        "07 固定工具目录：让大模型可调用图谱能力",
        "工具设计",
        [
            "get_person_labels：人物候选确认",
            "get_courtesy_name / get_art_name：查询字、号",
            "get_birth_death：查询生卒年",
            "get_teacher_relations：查询师承",
            "get_family_relations：查询亲属",
            "get_social_relations：查询交游",
            "get_school_membership / get_school_founder：查询流派",
            "get_pair_relations：查询两人关系",
        ],
        "课程对应",
        [
            "工具都有名称、参数和说明，符合课程中 tool calling 要求",
            "简单事实问题优先走工具，提高稳定性",
            "工具不足以回答时才转入 SPARQL 生成",
            "ToolNode 执行工具，ToolMessage 返回结果给模型",
        ],
        footer="文件依据：src/langchain_tools.py、src/tools.py",
    )

    make_two_col(
        prs.slides[15],
        "核心数据汇总：按实际项目口径更新",
        "图谱与本体",
        [
            "Person：5344",
            "Relation：12498",
            "本体类：9",
            "对象属性：18",
            "数据属性：7",
            "实际关系类型：15",
            "图谱三元组：约 17.36 万",
        ],
        "问答与工具",
        [
            "固定工具：13 个",
            "few-shot SPARQL：6 类典型问题",
            "LangGraph 节点：工具选择、ToolNode、SPARQL 生成、fallback",
            "前端功能：智能问答、高级 SPARQL、关系图、图分析、本体解释",
        ],
        footer="避免使用未核实数字：旧版 PPT 中的宣传口径已调整为当前仓库可验证口径。",
    )

    add_manual_ai_slide(prs)
    add_fewshot_slide(prs)
    add_demo_slide(prs)

    for slide in list(prs.slides):
        has_text = any(getattr(shape, "text", "").strip() for shape in slide.shapes)
        if not has_text:
            delete_slide(prs, slide)

    prs.save(OUT)
    print(OUT.resolve())


if __name__ == "__main__":
    main()
