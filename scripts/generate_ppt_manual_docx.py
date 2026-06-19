from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


OUT = Path("docs/ppt_manual_rules_and_ai_assistance_fixed.docx")


def set_font(style, size=None):
    style.font.name = "Microsoft YaHei"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        style.font.size = Pt(size)


def add_title(doc, text):
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run(text)
    run.bold = True
    run.font.size = Pt(18)


def add_bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_numbered(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Number")


def add_label(doc, text):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.bold = True


def add_code(doc, text):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    run.font.size = Pt(9.5)


def build_doc():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    for style_name in ["Normal", "Title", "Heading 1", "Heading 2", "Heading 3"]:
        set_font(doc.styles[style_name])
    doc.styles["Normal"].font.size = Pt(10.5)

    add_title(doc, "《印人传》知识图谱问答系统 PPT 制作思路")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("重点：5 分钟展示、课程技能点采分、人工设计规则与 AI 辅助边界").italic = True

    doc.add_heading("一、展示总原则", level=1)
    add_bullets(
        doc,
        [
            "展示以“功能展示 + 实现解释”为主，不做泛泛项目介绍。",
            "每一页都要对应课程评分点：知识抽取、RDF/Turtle、本体、实体链接、LangGraph、工具、few-shot SPARQL、前端。",
            "明确区分：哪些是人工设计，哪些是 AI 辅助，哪些是人工指导 AI 完成。",
            "不要说“全部自动完成”或“全量对齐完成”；对实体链接覆盖率要诚实说明为核心人物优先对齐。",
            "5 分钟内建议 7 页左右，每页只讲一个重点。",
        ],
    )

    doc.add_heading("二、5 分钟 PPT 页面结构", level=1)
    slides = [
        {
            "title": "第 1 页：项目目标与课程对应",
            "content": [
                "标题：《印人传》知识图谱问答系统。",
                "核心讲法：基于《印人传》构建知识图谱问答系统，完成知识抽取、RDF/Turtle 存储、实体链接、LangGraph 问答工作流、SPARQL 查询和网页前端展示。",
                "页面建议放总流程图：OCR/文本整理 -> 知识抽取 -> RDF/Turtle 图谱 -> ctext/cbdb 对齐 -> LangGraph 问答 -> 前端展示。",
            ],
            "manual": [
                "人工确定系统整体流程和课程技能点对应关系。",
                "人工整理项目模块：抽取、图谱、对齐、问答、前端。",
            ],
            "ai": ["AI 辅助知识抽取、复杂 SPARQL 生成和回答组织。"],
            "files": ["README.md", "BUGS_AND_TASKS.md"],
        },
        {
            "title": "第 2 页：知识抽取与人工 Schema 规则",
            "content": [
                "对应评分：知识抽取 30 分。",
                "重点讲法：这一部分不是让大模型自由生成，而是先人工定义抽取 Schema，再用 AI 辅助抽取。",
                "页面内容：实体类型、关系类型、输出格式。",
            ],
            "manual": [
                "人工定义实体类型：人物、地名、时间、字号、书体印风、流派、印章、作品。",
                "人工定义关系类型：父子、师承、交游、流派归属、开创、字号对应、籍贯、活动于、任职于、生于、卒于、擅长、创作、著有。",
                "人工规定输出格式：entities / relations / attributes，并要求 evidence 或 source_text 作为原文依据。",
            ],
            "ai": ["AI 根据人物传记文本抽取实体、属性、关系。"],
            "files": [
                "data/intermediate/extraction_schema.json",
                "docs/extraction/extraction_schema.md",
                "src/information_extraction.py",
            ],
        },
        {
            "title": "第 3 页：数据清洗与规则抽取",
            "content": [
                "这一页专门回答“哪些是你们设计的规则”。",
                "重点讲法：这些规则用于约束 AI 抽取结果，减少 OCR 或模型抽取带来的噪声。",
            ],
            "manual": [
                "去重规则：seen_entities 避免重复实体。",
                "停用词过滤：过滤“印章、书法、山水、花鸟、诸生、进士”等易误识别词。",
                "长度约束：过滤过短或过长的人名、字号、地名。",
                "中文字符过滤：减少非中文或 OCR 异常字符造成的噪声。",
                "原文窗口：使用 source_text / evidence 保留抽取依据，便于人工复核。",
                "关键词规则：字、号、别号、师、从XX学、与XX游、XX人、居XX、家于XX。",
            ],
            "ai": ["AI 抽取结果不足或 API 不可用时，规则抽取器作为补充。"],
            "files": ["scripts/extraction/rule_based_extract.py", "data/source/person_records_v5.json"],
        },
        {
            "title": "第 4 页：RDF / Turtle 与本体设计",
            "content": [
                "对应评分：RDF/Turtle 存储、本体构建附加项。",
                "页面建议放本体结构图：Person -> hasTeacher -> Person；Person -> belongsToSchool -> School；Person -> hasCourtesyName -> CourtesyName。",
                "数据可放：Person 5344，Relation 12498，本体类 9，对象属性 18，数据属性 7，实际关系类型 15。",
            ],
            "manual": [
                "人工设计 9 个核心类：Person、Place、TimePeriod、CourtesyName、CalligraphyStyle、School、Seal、Work、Relation。",
                "人工设计人物关系与属性论元：hasTeacher、fatherOf、hasFriend、belongsToSchool、foundsSchool、hasCourtesyName、hasArtName、birthPlace、skilledIn、authored 等。",
                "人工确定 domain/range，例如 hasTeacher: Person -> Person，belongsToSchool: Person -> School。",
                "人工决定采用 RDF/Turtle 作为正式图谱格式。",
            ],
            "ai": ["AI 可辅助从抽取结果中形成事实，但本体类和属性设计需要人工确定。"],
            "files": ["data/kg/schema.ttl", "data/kg/core.ttl", "data/kg/aligned.ttl", "src/graph_store.py"],
        },
        {
            "title": "第 5 页：实体链接与知识补充",
            "content": [
                "对应评分：实体链接和知识补充 20 分。",
                "重点讲法：当前完成核心人物优先对齐，不说全量完成。",
                "当前数据：20 个唯一人物，104 个 aligned 资源，165 条 owl:sameAs。",
            ],
            "manual": [
                "人工消歧规则：字号完全相同 + 籍贯相同为高置信。",
                "人工消歧规则：字号相同 + 籍贯同区域为中置信。",
                "人工消歧规则：仅字号相同为低置信，需要人工确认。",
                "人工消歧规则：朝代、生卒年不匹配时排除候选。",
                "多候选规则：优先籍贯完全匹配，其次同省同府。",
                "外部知识补充：owl:sameAs、cbdbId、ctextId、bornIn、diedIn、alignmentStatus、alignmentMethod。",
            ],
            "ai": [
                "后续计划加入大模型打分消歧：在规则无法唯一判断时，让模型输出候选、分数和理由。当前项目中 LLM 打分消歧实现未找到，不要说已经完成。"
            ],
            "files": ["data/kg/alignment_rules.md", "docs/kg_alignment/alignment_rules.md", "data/kg/aligned.ttl"],
        },
        {
            "title": "第 6 页：LangGraph 问答工作流与工具设计",
            "content": [
                "对应评分：问答系统 30 分，包括 LangGraph 工作流、工具调用、few-shot SPARQL。",
                "页面建议放流程图：用户问题 -> LLM 判断工具 -> 固定工具查询 -> 有结果回答；无结果 -> few-shot 生成 SPARQL -> 执行；失败 -> fallback。",
            ],
            "manual": [
                "人工设计工具目录：get_courtesy_name、get_art_name、get_birth_death、get_teacher_relations、get_family_relations、get_social_relations、get_school_membership、get_pair_relations。",
                "人工设计路由规则：字、号、生卒年、师承、亲属、交游、流派优先走固定工具。",
                "人工设计 fallback 顺序：工具失败 -> SPARQL 失败 -> 大模型谨慎回答。",
                "人工设计 few-shot SPARQL 示例：典型自然语言问题对应可执行 SPARQL 模板。",
                "人工设计别名归一：文征明/文徵明/征仲，吴门印派/吴门派/吴门。",
            ],
            "ai": ["AI 辅助判断是否调用工具，辅助生成复杂 SPARQL，辅助根据查询结果组织自然语言回答。"],
            "files": [
                "src/tool_calling_workflow.py",
                "src/langchain_tools.py",
                "src/tools.py",
                "src/fewshot_sparql.py",
                "data/examples/fewshot_sparql.md",
                "src/workflow.py",
            ],
        },
        {
            "title": "第 7 页：网页功能展示与总结",
            "content": [
                "这一页直接功能演示，不放太多文字。",
                "展示顺序：智能问答 -> 高级 SPARQL -> 关系网络图 -> 图结构分析 -> 本体解释。",
                "建议演示问题：文彭的字和号是什么？文徵明与文彭是什么关系？谁开创了吴门印派？",
            ],
            "manual": [
                "人工设计前端功能组织：智能问答、高级 SPARQL、关系网络图、图结构分析、本体解释、示例问题。",
                "人工设计图分析口径：中心性、社区、路径、流派分析基于人物关系图谱。",
                "人工设计本体可视化解释：点击关系边展示实例关系、本体属性、domain、range、说明。",
            ],
            "ai": ["AI 辅助问答生成和部分代码实现；前端模块组织与演示路线由人工确定。"],
            "files": ["templates/index.html", "static/styles.css", "src/graph_analysis.py", "src/web.py"],
        },
    ]

    for slide in slides:
        doc.add_heading(slide["title"], level=2)
        add_label(doc, "页面内容 / 讲什么：")
        add_bullets(doc, slide["content"])
        add_label(doc, "人工设计或人工指导完成：")
        add_bullets(doc, slide["manual"])
        add_label(doc, "AI 辅助部分：")
        add_bullets(doc, slide["ai"])
        add_label(doc, "可引用文件依据：")
        add_bullets(doc, slide["files"])

    doc.add_heading("三、可重点强调的人工完成内容清单", level=1)
    add_bullets(
        doc,
        [
            "抽取 Schema 设计：实体类型、关系类型、输出结构。",
            "数据清洗规则：去重、停用词、长度限制、中文字符过滤、原文证据保留。",
            "关系抽取规则：通过正则和关键词识别字号、籍贯、师承、父子、交游、流派等。",
            "RDF/Turtle 与本体设计：9 个类、18 个对象属性、7 个数据属性、domain/range。",
            "实体链接消歧规则：字号、籍贯、朝代、生卒年、多候选优先级。",
            "工具设计规则：为大模型提供可调用工具及参数说明。",
            "LangGraph 路由规则：固定工具优先，复杂问题走 SPARQL，失败进入 fallback。",
            "few-shot SPARQL 示例设计：人工准备典型问句和 SPARQL 模板。",
            "人名/流派别名归一：解决文征明/文徵明、吴门/吴门印派等表达差异。",
            "图结构分析口径：中心性、社区、路径、流派分析基于人物关系网络。",
            "前端展示规则：功能分区、示例问题、本体解释展示、关系图交互。",
        ],
    )

    doc.add_heading("四、需要谨慎表述的地方", level=1)
    add_bullets(
        doc,
        [
            "OCR：仓库中保留了清洗后的 JSON 数据，但未找到 OCR 脚本。可说“PDF/OCR 后文本已整理为人物记录 JSON”，不要说 OCR 流程完整保存在仓库。",
            "JSON 转 TTL：仓库中有 schema.ttl、core.ttl、aligned.ttl 和 RDFLib 加载查询代码，但未找到 JSON 自动转换 TTL 的脚本。可说“保留了最终 RDF/Turtle 图谱文件”。",
            "实体链接：不要说全量对齐完成。建议说“核心人物优先对齐，当前 20 个唯一人物、104 个 aligned 资源、165 条 owl:sameAs”。",
            "LLM 打分消歧：当前未找到实现代码，应说“后续计划加入”或“作为待补充改进”，不要说已完成。",
        ],
    )

    doc.add_heading("五、5 分钟时间分配", level=1)
    add_bullets(
        doc,
        [
            "0:00-0:30 项目目标和总流程。",
            "0:30-1:20 知识抽取 + 数据清洗规则。",
            "1:20-2:00 RDF/Turtle + 本体设计。",
            "2:00-2:40 实体链接 + 消歧规则。",
            "2:40-3:40 LangGraph 问答工作流 + 工具 + few-shot SPARQL。",
            "3:40-4:40 前端功能演示。",
            "4:40-5:00 总结人工设计与 AI 辅助边界。",
        ],
    )

    doc.add_heading("六、推荐总结话术", level=1)
    add_label(doc, "可以放在最后一页或答辩结尾：")
    add_code(
        doc,
        "本项目不是单纯调用大模型生成答案，而是用人工设计的 Schema、清洗规则、本体、实体链接消歧规则、工具目录、few-shot SPARQL 示例和 LangGraph 路由规则约束 AI，使大模型服务于本地 RDF 知识图谱查询。",
    )

    doc.add_heading("七、PPT 可直接使用的页标题", level=1)
    add_numbered(
        doc,
        [
            "项目目标与课程技能点对应",
            "知识抽取：人工 Schema + AI 辅助抽取",
            "数据清洗与规则抽取",
            "RDF/Turtle 与本体设计",
            "实体链接、规则消歧与知识补充",
            "LangGraph 问答工作流：工具 + SPARQL + fallback",
            "网页前端功能展示与总结",
        ],
    )
    return doc


if __name__ == "__main__":
    OUT.parent.mkdir(parents=True, exist_ok=True)
    build_doc().save(OUT)
    print(OUT.resolve())
